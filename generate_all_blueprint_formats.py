import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, HRFlowable
)
from reportlab.pdfgen import canvas

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def build_word_blueprint(filename="docs/integrieval-blueprint.docx"):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("IntegriEval — Blueprint de Arquitectura (Estándar Innovares) | Septiembre 2026")
        hrun.font.name = "Arial"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # Title
    p_t = doc.add_paragraph()
    p_t.paragraph_format.space_before = Pt(20)
    p_t.paragraph_format.space_after = Pt(4)
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_t.add_run("IntegriEval — Blueprint de Arquitectura")
    r_t.bold = True
    r_t.font.name = "Arial"
    r_t.font.size = Pt(24)
    r_t.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    p_s = doc.add_paragraph()
    p_s.paragraph_format.space_after = Pt(20)
    p_s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_s = p_s.add_run("Plataforma Semi-Automatizada de Evaluación de Integridad Académica y Defensas Flash\nEstándar El Arquitecto Innovares — Trabajo de Título (INFO1197)")
    r_s.font.name = "Arial"
    r_s.font.size = Pt(11)
    r_s.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    # Metadata
    tbl_m = doc.add_table(rows=3, cols=2)
    tbl_m.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_m.autofit = False
    m_data = [
        [("Autores:", True), ("Sebastian Cisternas, Benjamin Sobarzo", False), ("Versión:", True), ("1.0.0 (Release para Validación)", False)],
        [("Institución:", True), ("Facultad de Ingeniería Civil Informática", False), ("Fecha:", True), ("Septiembre 2026", False)],
        [("Repositorio:", True), ("github.com/blosttt/IntegriEval", False), ("Hosting:", True), ("Vercel + GMKtec M5 + Cloudflare Tunnel", False)]
    ]
    for r_i, row in enumerate(m_data):
        c1 = tbl_m.cell(r_i, 0)
        c2 = tbl_m.cell(r_i, 1)
        c1.width = Inches(3.3)
        c2.width = Inches(3.3)
        set_cell_background(c1, "F8FAFC")
        set_cell_background(c2, "F8FAFC")
        set_cell_margins(c1, 60, 60, 100, 100)
        set_cell_margins(c2, 60, 60, 100, 100)
        
        p1 = c1.paragraphs[0]
        r1a = p1.add_run(row[0][0] + " ")
        r1a.bold = True
        r1a.font.size = Pt(9)
        r1b = p1.add_run(row[1][0])
        r1b.font.size = Pt(9)
        
        p2 = c2.paragraphs[0]
        r2a = p2.add_run(row[2][0] + " ")
        r2a.bold = True
        r2a.font.size = Pt(9)
        r2b = p2.add_run(row[3][0])
        r2b.font.size = Pt(9)

    set_table_borders(tbl_m, color="93C5FD", sz="4", val="single")

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    def add_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    add_h1("1. Visión y Propósito del Proyecto")
    add_p("IntegriEval resuelve la crisis de evaluación en la educación superior provocada por los LLMs (GPT-4, Claude). En lugar de usar detectores binarios de 'caja negra' propensos a falsos positivos, implementa un modelo híbrido: extracción estructurada a Markdown, análisis cualitativo contextualizado con el syllabus del curso, pruebas flash interactivas y citación focalizada a defensas orales presenciales.")

    add_h1("2. Stack Tecnológico y Arquitectura de Hosting")
    add_p("• Frontend: Next.js 15 (App Router) + Tailwind CSS v4 + shadcn/ui alojado en Vercel Edge.\n• Backend: FastAPI (Python 3.14) + WebSockets sobre GMKtec M5 Ultra 24/7.\n• Base de Datos: SQLite local / PostgreSQL (Supabase / NAS) con SQLAlchemy 2.0.\n• Inferencia IA: Claude 3.5 Sonnet con fallback transparente a Mock Engine.\n• Exposición: Cloudflare Tunnel nombrado con HTTPS fija (api.integrieval.innovares.cl).")

    add_h1("3. Service Blueprint y Flujo Operativo")
    add_p("1. Ingesta: El docente carga nómina CSV y lote de PDFs (extracción semántica con pymupdf4llm).\n2. Análisis: La IA evalúa cumplimiento de pauta y probabilidad de IA con nota preliminar.\n3. Supervisión: El docente edita el pool de 20 preguntas y aprueba exactamente 10 reactivos.\n4. Flash Test: El alumno rinde la prueba en tiempo real vía WebSocket (60s/pregunta) sin crear cuenta.\n5. Ruteo y Cierre: Se agendan citas a casos <50%, >=95% o 10% aleatorio, registrando la nota final en el AuditLog inmutable.")

    add_h1("4. Checklist de Prelanzamiento (P0 / P1 / P2)")
    add_p("• P0 (Bloqueantes): Tokens UUIDv4 con vigencia 48h, temporizador server-side de 60s, hash bcrypt en docentes y AuditLog inmutable.\n• P1 (Importantes): Fallback automático a Mock Engine, reconexión de WebSocket <= 120s y sanitización CSV.\n• P2 (Deseables): Integración LTI con LMS Moodle/Canvas y exportación de reportes ejecutivos.")

    doc.save(filename)
    print("Word Blueprint build complete:", filename)


class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        if self._pageNumber == 1:
            return
        self.saveState()
        self.setFont("Helvetica-Bold", 8)
        self.setFillColor(colors.HexColor("#475569"))
        self.drawString(54, 750, "IntegriEval — Blueprint de Arquitectura (Estándar Innovares)")
        self.drawRightString(558, 750, "Versión 1.0 (Septiembre 2026)")
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        self.line(54, 742, 558, 742)
        self.line(54, 45, 558, 45)
        self.setFont("Helvetica", 8)
        self.drawString(54, 32, "Trabajo de Título (INFO1197) — OTEC Innovares / UCT")
        self.drawRightString(558, 32, f"Página {self._pageNumber} de {page_count}")
        self.restoreState()


def build_pdf_blueprint(filename="docs/integrieval-blueprint.pdf"):
    doc = SimpleDocTemplate(filename, pagesize=letter, leftMargin=54, rightMargin=54, topMargin=64, bottomMargin=54)
    styles = getSampleStyleSheet()
    
    c_primary = colors.HexColor("#1E3A8A")
    c_secondary = colors.HexColor("#2563EB")
    c_indigo = colors.HexColor("#4338CA")
    c_text = colors.HexColor("#0F172A")
    c_muted = colors.HexColor("#475569")
    c_bg = colors.HexColor("#F8FAFC")
    
    styles.add(ParagraphStyle('CovT', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=24, leading=28, textColor=c_primary, alignment=1))
    styles.add(ParagraphStyle('CovS', parent=styles['Normal'], fontName='Helvetica', fontSize=11, leading=15, textColor=c_indigo, alignment=1))
    styles.add(ParagraphStyle('SecH', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=12, leading=15, textColor=c_primary, spaceBefore=10, spaceAfter=4, keepWithNext=True))
    styles.add(ParagraphStyle('SubH', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=9.5, leading=13, textColor=c_secondary, spaceBefore=6, spaceAfter=2, keepWithNext=True))
    styles.add(ParagraphStyle('Body', parent=styles['Normal'], fontName='Helvetica', fontSize=8, leading=11, textColor=c_text, spaceAfter=3))
    styles.add(ParagraphStyle('TCell', parent=styles['Normal'], fontName='Helvetica', fontSize=7.5, leading=10, textColor=c_text))
    styles.add(ParagraphStyle('THead', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=7.5, leading=10, textColor=colors.white))
    styles.add(ParagraphStyle('BoxT', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=8.5, leading=11, textColor=c_indigo))

    story = []

    # Portada
    story.append(Spacer(1, 30))
    story.append(Paragraph("IntegriEval", styles['CovT']))
    story.append(Spacer(1, 6))
    story.append(Paragraph("Blueprint de Arquitectura de Software y Service Blueprint<br/>Estándar Metodológico El Arquitecto Innovares", styles['CovS']))
    story.append(Spacer(1, 15))
    story.append(HRFlowable(width="80%", thickness=1.5, color=c_secondary, spaceBefore=8, spaceAfter=15))

    meta = [
        [Paragraph("<b>Proyecto:</b> Trabajo de Título (INFO1197) — IntegriEval", styles['TCell']), Paragraph("<b>Autores:</b> Sebastian Cisternas, Benjamin Sobarzo", styles['TCell'])],
        [Paragraph("<b>Versión:</b> 1.0.0 (Release Septiembre 2026)", styles['TCell']), Paragraph("<b>Carrera:</b> Ingeniería Civil Informática", styles['TCell'])],
        [Paragraph("<b>Stack:</b> Next.js 15 + FastAPI + Claude 3.5 + SQLite/Postgres", styles['TCell']), Paragraph("<b>Hosting:</b> Vercel + GMKtec M5 + Cloudflare Tunnel", styles['TCell'])]
    ]
    t_m = Table(meta, colWidths=[250, 254])
    t_m.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), c_bg),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#CBD5E1")),
        ('PADDING', (0,0), (-1,-1), 5),
    ]))
    story.append(t_m)
    story.append(Spacer(1, 15))

    # Sección 1
    story.append(Paragraph("1. Visión y Propósito del Sistema", styles['SecH']))
    story.append(Paragraph("IntegriEval es una solución semi-automatizada que asiste a docentes universitarios en la evaluación de informes académicos frente al avance de la IA generativa. En vez de recurrir a detectores binarios tradicionales, implementa un modelo de <b>verificación de aprendizaje auténtico</b> con extracción semántica a Markdown, generación contextual de reactivos y defensas orales focalizadas con control 100% docente (<i>Human-in-the-Loop</i>).", styles['Body']))

    # Sección 2
    story.append(Paragraph("2. Service Blueprint (5 Niveles de Operación)", styles['SecH']))
    bp_data = [
        [Paragraph("<b>Nivel</b>", styles['THead']), Paragraph("<b>1. Ingesta CSV/PDF</b>", styles['THead']), Paragraph("<b>2. Análisis IA</b>", styles['THead']), Paragraph("<b>3. Pool Reactivos</b>", styles['THead']), Paragraph("<b>4. Flash Test WS</b>", styles['THead']), Paragraph("<b>5. Cierre & Cita</b>", styles['THead'])],
        [Paragraph("<b>Frontstage</b>", styles['TCell']), Paragraph("Docente sube nómina y PDFs.", styles['TCell']), Paragraph("Docente revisa nota y % IA.", styles['TCell']), Paragraph("Docente aprueba 10 preguntas.", styles['TCell']), Paragraph("Alumno rinde test en vivo.", styles['TCell']), Paragraph("Docente interroga y cierra acta.", styles['TCell'])],
        [Paragraph("<b>Frontend</b>", styles['TCell']), Paragraph("Dropzone y parser CSV reactivo.", styles['TCell']), Paragraph("Tarjetas de rúbrica y feedback.", styles['TCell']), Paragraph("Editor con validación N=10.", styles['TCell']), Paragraph("WebSocket con reloj de 60s.", styles['TCell']), Paragraph("Panel de actas con AuditLog.", styles['TCell'])],
        [Paragraph("<b>Backstage</b>", styles['TCell']), Paragraph("POST /api/students /reports", styles['TCell']), Paragraph("Worker asíncrono FastAPI.", styles['TCell']), Paragraph("Generador UUIDv4 + SMTP.", styles['TCell']), Paragraph("WS /ws/flash/{token}", styles['TCell']), Paragraph("POST /api/appointments/close", styles['TCell'])],
        [Paragraph("<b>Soporte</b>", styles['TCell']), Paragraph("pymupdf4llm (Markdown).", styles['TCell']), Paragraph("Claude 3.5 / MockEngine.", styles['TCell']), Paragraph("Gmail SMTP (aiosmtplib).", styles['TCell']), Paragraph("Temporizador server-side.", styles['TCell']), Paragraph("AuditLog inmutable en BD.", styles['TCell'])],
    ]
    t_bp = Table(bp_data, colWidths=[64, 88, 88, 88, 88, 88])
    t_bp.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), c_primary),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, c_bg]),
        ('PADDING', (0,0), (-1,-1), 3),
    ]))
    story.append(t_bp)
    story.append(Spacer(1, 8))

    # Sección 3
    story.append(Paragraph("3. Tech Stack y Decisión de Hosting (Home Lab)", styles['SecH']))
    h_data = [
        [Paragraph("<b>Componente</b>", styles['THead']), Paragraph("<b>Tecnología</b>", styles['THead']), Paragraph("<b>Destino de Hosting</b>", styles['THead']), Paragraph("<b>Justificación Operativa</b>", styles['THead'])],
        [Paragraph("<b>Frontend SPA</b>", styles['TCell']), Paragraph("Next.js 15 + Tailwind CSS v4", styles['TCell']), Paragraph("Vercel Edge Network", styles['TCell']), Paragraph("Baja latencia global, SSL automático y cero costo de hosting.", styles['TCell'])],
        [Paragraph("<b>API & WebSockets</b>", styles['TCell']), Paragraph("FastAPI (Python 3.14)", styles['TCell']), Paragraph("GMKtec M5 Ultra (24/7)", styles['TCell']), Paragraph("Nodo headless de bajo consumo; soporta 100 WS concurrentes.", styles['TCell'])],
        [Paragraph("<b>Base de Datos</b>", styles['TCell']), Paragraph("PostgreSQL / SQLite + SQLAlchemy", styles['TCell']), Paragraph("Supabase / NAS UGREEN", styles['TCell']), Paragraph("Transacciones ACID seguras y almacenamiento de AuditLog.", styles['TCell'])],
        [Paragraph("<b>Túnel Seguro</b>", styles['TCell']), Paragraph("Cloudflare Tunnel", styles['TCell']), Paragraph("Nombrado (HTTPS)", styles['TCell']), Paragraph("Expone el backend sin abrir puertos ni IPs públicas fijas.", styles['TCell'])],
    ]
    t_h = Table(h_data, colWidths=[80, 120, 110, 194])
    t_h.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), c_primary),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, c_bg]),
        ('PADDING', (0,0), (-1,-1), 3.5),
    ]))
    story.append(t_h)
    story.append(Spacer(1, 8))

    # Sección 4
    story.append(Paragraph("4. Checklist de Prelanzamiento y Reglas No Negociables", styles['SecH']))
    story.append(Paragraph("<b>• P0 (Bloqueantes):</b> Tokens UUIDv4 con 48h de vigencia, temporizador en servidor (60s), contraseñas con bcrypt ($\ge 12$) y AuditLog inmutable.<br/><b>• P1 (Importantes):</b> Fallback automático a Mock Engine ante cortes de API, reconexión de WebSockets $\le 120$s y cumplimiento de la Ley N° 21.719.<br/><b>• P2 (Deseables):</b> Integración mediante LTI con LMS Moodle / Canvas y exportación ejecutiva de actas.", styles['Body']))

    doc.build(story, canvasmaker=NumberedCanvas)
    print("PDF Blueprint build complete:", filename)


if __name__ == "__main__":
    os.makedirs("docs", exist_ok=True)
    build_word_blueprint()
    build_pdf_blueprint()
